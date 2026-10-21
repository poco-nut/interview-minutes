#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 minutes_data.json 渲染成《XX访谈纪要》docx。

版式源自复星创富公司纪要样本（卓然聚变 / 星阑科技 / 云驰半导体三份互证）：
  标题      居中 加粗 16pt 段后 12pt
  元信息    12pt 段后 3pt
  分节标题  加粗 12pt 段后 3pt
  子节标题  加粗 12pt 缩进一级 段后 3pt
  引导短语  段内加粗 run + 全角冒号，后接正文
  Q&A 问题  加粗 12pt 段后 3pt
  「答：」  加粗 12pt 段后 3pt
  正文      12pt 段后 3pt 行距 1.15
  字体      中文 宋体 / 西文 Times New Roman

用法：
  PYTHONIOENCODING=utf-8 python render_minutes.py --data minutes_data.json
  PYTHONIOENCODING=utf-8 python render_minutes.py --data minutes_data.json --out "XX访谈纪要20260920.docx"
"""

import argparse
import datetime
import json
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

CN_FONT = "宋体"
EN_FONT = "Times New Roman"

BODY_SIZE = Pt(12)
TITLE_SIZE = Pt(16)

# w:after="60" -> 3pt；标题 w:after="240" -> 12pt
AFTER_BODY = Pt(3)
AFTER_TITLE = Pt(12)

LINE_SPACING = 1.15
SUB_INDENT = Pt(21)  # 约两个字符，对应样本的缩进一级

CN_NUMERALS = "一二三四五六七八九十"


def set_font(run, size=BODY_SIZE, bold=False):
    """给 run 同时设定中西文字体，否则中文会回落成 Calibri。"""
    run.font.size = size
    run.font.bold = bold
    run.font.name = EN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)
    rfonts.set(qn("w:cs"), EN_FONT)
    rfonts.set(qn("w:eastAsia"), CN_FONT)


def add_para(doc, after=AFTER_BODY, align=None, indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = after
    pf.line_spacing = LINE_SPACING
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = indent
    return p


def add_plain(doc, text, bold=False, after=AFTER_BODY, align=None, indent=None, size=BODY_SIZE):
    p = add_para(doc, after=after, align=align, indent=indent)
    set_font(p.add_run(text), size=size, bold=bold)
    return p


def add_lead_para(doc, lead, text):
    """引导短语段：加粗短语 + 全角冒号 + 常规正文，同一段内。"""
    p = add_para(doc)
    if lead:
        set_font(p.add_run(lead.rstrip("：:") + "："), bold=True)
    set_font(p.add_run(text))
    return p


def add_para_block(doc, paras, indent=None):
    """渲染段落块。每项可以是字符串，或 {"lead": ..., "text": ...}。"""
    for item in paras or []:
        if isinstance(item, str):
            add_plain(doc, item, indent=indent)
        else:
            add_lead_para_indented(doc, item.get("lead"), item.get("text", ""), indent)


def add_lead_para_indented(doc, lead, text, indent=None):
    p = add_para(doc, indent=indent)
    if lead:
        set_font(p.add_run(lead.rstrip("：:") + "："), bold=True)
    set_font(p.add_run(text))
    return p


def build(data, out_path):
    doc = Document()

    # 文档级默认字体，兜住任何未显式设格式的文字
    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    # ---- 标题 ----
    add_plain(
        doc,
        data["title"],
        bold=True,
        after=AFTER_TITLE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=TITLE_SIZE,
    )

    # ---- 元信息 ----
    for item in data.get("meta", []):
        p = add_para(doc)
        set_font(p.add_run(item["label"] + "："))
        set_font(p.add_run(item["value"]))

    # ---- 项目简介 ----
    add_plain(doc, data.get("intro_title", ""), bold=True)
    add_plain(doc, data["intro"])

    # ---- 分节正文 ----
    for sec in data["sections"]:
        add_plain(doc, sec["heading"], bold=True)
        add_para_block(doc, sec.get("paras"))
        for sub in sec.get("subsections", []):
            add_plain(doc, sub["heading"], bold=True, indent=SUB_INDENT)
            add_para_block(doc, sub.get("paras"), indent=SUB_INDENT)
        add_para_block(doc, sec.get("outro"))

    # ---- Q&A ----
    qa = data.get("qa")
    if qa:
        add_plain(doc, qa.get("heading", "Q&A"), bold=True)
        for it in qa.get("items", []):
            asker = it.get("asker")
            head = it["no"] + ("（%s）" % asker if asker else "") + "：" + it["question"]
            add_plain(doc, head, bold=True)
            add_plain(doc, "答：", bold=True)
            answer = it.get("answer", "")
            for para in (answer if isinstance(answer, list) else [answer]):
                add_plain(doc, para)

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="渲染《XX访谈纪要》docx")
    ap.add_argument("--data", required=True, help="minutes_data.json 路径")
    ap.add_argument("--out", help="输出文件名，省则按 <title><今天>.docx")
    args = ap.parse_args()

    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)

    for key in ("title", "intro", "sections"):
        if not data.get(key):
            sys.exit("minutes_data.json 缺少必填字段：%s" % key)

    out = args.out or "%s%s.docx" % (data["title"], datetime.date.today().strftime("%Y%m%d"))
    build(data, out)
    print("已生成：%s" % out)


if __name__ == "__main__":
    main()
